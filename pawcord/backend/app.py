import base64
import csv
import io
import json
import os
import re
from datetime import datetime
from typing import Literal

import openpyxl
from anthropic import Anthropic
from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from supabase import create_client

load_dotenv()

app = FastAPI(title="Pawcord API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

supabase = create_client(
    os.environ["SUPABASE_URL"],
    os.environ["SUPABASE_SERVICE_KEY"],
)

anthropic = Anthropic(api_key=os.environ["ANTHROPIC_API_KEY"])

TABLE = "checkup_records"

PARSE_SYSTEM_PROMPT = """You are a veterinary lab report parser. Extract blood test values from the provided document.
Return ONLY a valid JSON object with this exact schema — no markdown, no explanation:
{
  "exam_date": "YYYY-MM-DD or null",
  "hospital": "string or null",
  "cat_name": "string or null",
  "weight_kg": number or null,
  "cbc": {
    "RBC":  {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "HGB":  {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "HCT":  {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "MCV":  {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "MCH":  {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "MCHC": {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "WBC":  {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "NEU":  {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "LYM":  {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "MONO": {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "EOS":  {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "PLT":  {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null}
  },
  "chemistry": {
    "BUN":  {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "CREA": {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "ALT":  {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "AST":  {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "ALP":  {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "GGT":  {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "TP":   {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "ALB":  {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "GLOB": {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "GLU":  {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "CHOL": {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "PHOS": {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "Ca":   {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "Na":   {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "K":    {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null},
    "Cl":   {"value": number, "unit": "string", "ref_low": number or null, "ref_high": number or null}
  }
}
If a marker is not present in the document, set its value to null (not the object, just null).
If a field within a marker object is not available, set that field to null."""


def _detect_file_type(
    filename: str, content_type: str
) -> Literal["pdf", "image", "excel", "csv"]:
    name = filename.lower()
    if name.endswith(".pdf") or content_type == "application/pdf":
        return "pdf"
    if name.endswith((".jpg", ".jpeg")) or content_type in ("image/jpeg", "image/jpg"):
        return "image"
    if name.endswith(".png") or content_type == "image/png":
        return "image"
    if name.endswith((".xlsx", ".xls")) or "spreadsheet" in content_type:
        return "excel"
    if name.endswith(".csv") or content_type in ("text/csv", "application/csv"):
        return "csv"
    raise HTTPException(status_code=415, detail=f"Unsupported file type: {filename}")


def _build_claude_message(
    file_type: Literal["pdf", "image", "excel", "csv"],
    raw_bytes: bytes,
    filename: str,
) -> list[dict]:
    if file_type == "pdf":
        return [
            {
                "type": "document",
                "source": {
                    "type": "base64",
                    "media_type": "application/pdf",
                    "data": base64.b64encode(raw_bytes).decode(),
                },
            },
            {"type": "text", "text": "Extract all blood test values from this veterinary lab report."},
        ]

    if file_type == "image":
        # Determine media type
        media_type = "image/jpeg"
        if filename.lower().endswith(".png"):
            media_type = "image/png"
        return [
            {
                "type": "image",
                "source": {
                    "type": "base64",
                    "media_type": media_type,
                    "data": base64.b64encode(raw_bytes).decode(),
                },
            },
            {"type": "text", "text": "Extract all blood test values from this veterinary lab report image."},
        ]

    if file_type == "excel":
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
        lines = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            lines.append(f"## Sheet: {sheet_name}")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                lines.append(" | ".join(cells))
        text = "\n".join(lines)
        return [
            {"type": "text", "text": f"Lab report data from Excel file:\n\n{text}\n\nExtract all blood test values from this veterinary lab report."},
        ]

    # csv
    text = raw_bytes.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    lines = [" | ".join(row) for row in reader]
    table_text = "\n".join(lines)
    return [
        {"type": "text", "text": f"Lab report data from CSV:\n\n{table_text}\n\nExtract all blood test values from this veterinary lab report."},
    ]


def _parse_with_claude(content_blocks: list[dict]) -> dict:
    last_error = None
    for attempt in range(3):
        response = anthropic.messages.create(
            model="claude-opus-4-6",
            max_tokens=4096,
            system=PARSE_SYSTEM_PROMPT,
            messages=[{"role": "user", "content": content_blocks}],
        )
        raw = response.content[0].text.strip()
        # Strip markdown code fences if present
        raw = re.sub(r"^```(?:json)?\n?", "", raw)
        raw = re.sub(r"\n?```$", "", raw)
        try:
            return json.loads(raw)
        except json.JSONDecodeError as e:
            last_error = e
            continue
    raise HTTPException(
        status_code=422,
        detail={"message": "AI parsing failed to return valid JSON", "error": str(last_error)},
    )


def _build_analysis_prompt(records: list[dict]) -> str:
    if not records:
        return "No records available for analysis."

    lines = ["Veterinary blood test history for analysis:\n"]
    for rec in sorted(records, key=lambda r: r.get("exam_date") or ""):
        date = rec.get("exam_date", "unknown date")
        cat = rec.get("cat_name") or "cat"
        lines.append(f"\n--- {date} ({cat}) ---")
        for panel in ("cbc", "chemistry"):
            panel_data = rec.get(panel) or {}
            for marker, item in panel_data.items():
                if item is None:
                    continue
                v = item.get("value")
                u = item.get("unit", "")
                rl = item.get("ref_low")
                rh = item.get("ref_high")
                status = "normal"
                if rl is not None and v < rl:
                    status = "LOW"
                elif rh is not None and v > rh:
                    status = "HIGH"
                lines.append(f"  {marker}: {v} {u} [{status}] (ref {rl}–{rh})")

    prompt_text = "\n".join(lines)
    prompt_text += (
        "\n\nPlease provide a concise veterinary-style summary of trends, notable values, "
        "and any concerns worth discussing with a vet. Write in plain language suitable for a cat owner. "
        "Keep it under 300 words."
    )
    return prompt_text


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------


@app.get("/health")
def health():
    return {"status": "ok"}


@app.post("/api/upload", status_code=201)
async def upload(file: UploadFile = File(...)):
    raw_bytes = await file.read()

    if len(raw_bytes) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="File too large (max 10MB)")

    filename = file.filename or "upload"
    content_type = file.content_type or ""

    file_type = _detect_file_type(filename, content_type)
    content_blocks = _build_claude_message(file_type, raw_bytes, filename)
    parsed = _parse_with_claude(content_blocks)

    record_id = datetime.now().strftime("%Y%m%d%H%M%S")

    record = {
        "id": record_id,
        "exam_date": parsed.get("exam_date"),
        "hospital": parsed.get("hospital"),
        "cat_name": parsed.get("cat_name"),
        "weight_kg": parsed.get("weight_kg"),
        "filename": filename,
        "file_type": file_type,
        "cbc": parsed.get("cbc") or {},
        "chemistry": parsed.get("chemistry") or {},
    }

    result = supabase.table(TABLE).insert(record).execute()
    return result.data[0]


@app.get("/api/records")
def get_records():
    result = supabase.table(TABLE).select("*").order("exam_date", desc=True).execute()
    return result.data


@app.get("/api/records/{record_id}")
def get_record(record_id: str):
    result = supabase.table(TABLE).select("*").eq("id", record_id).execute()
    if not result.data:
        raise HTTPException(status_code=404, detail="Record not found")
    return result.data[0]


@app.delete("/api/records/{record_id}", status_code=204)
def delete_record(record_id: str):
    result = supabase.table(TABLE).select("id").eq("id", record_id).execute()
    if not result.data:
        raise HTTPException(status_code=404, detail="Record not found")
    supabase.table(TABLE).delete().eq("id", record_id).execute()


@app.get("/api/analyze")
def analyze():
    result = supabase.table(TABLE).select("*").order("exam_date", desc=False).execute()
    records = result.data

    if not records:
        return {"summary": "No records found. Upload a blood test report to get started."}

    prompt_text = _build_analysis_prompt(records)

    response = anthropic.messages.create(
        model="claude-opus-4-6",
        max_tokens=1024,
        messages=[{"role": "user", "content": prompt_text}],
    )
    return {"summary": response.content[0].text.strip()}
